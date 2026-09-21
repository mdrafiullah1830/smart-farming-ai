from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

OUT = "/Users/mdrafiullah/smart_farming_ai/SAU_Agricultural_Data_Equipment_Application.docx"

doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.27)
sec.page_height = Inches(11.69)
sec.top_margin = Inches(0.65)
sec.bottom_margin = Inches(0.65)
sec.left_margin = Inches(0.8)
sec.right_margin = Inches(0.8)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Times New Roman"
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(4)
normal.paragraph_format.line_spacing = 1.05

def add(text="", bold=False, align=None, after=4, before=0):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(before)
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(11)
    r.bold = bold
    return p

add("Date: ____ / ____ / 2026", align=WD_ALIGN_PARAGRAPH.RIGHT, after=6)
add("To", after=0)
add("The Director", after=0)
add("Sher-e-Bangla Agricultural University Research System (SAURES)", after=0)
add("Sher-e-Bangla Agricultural University, Dhaka-1207", after=7)

p = add(after=7)
r = p.add_run("Subject: ")
r.bold = True; r.font.name = "Times New Roman"; r.font.size = Pt(11)
r = p.add_run("Request for Agricultural Datasets and Access to Soil-Testing Equipment")
r.bold = True; r.font.name = "Times New Roman"; r.font.size = Pt(11)

add("Respected Sir/Madam,", after=5)

p = add(after=5)
p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
p.add_run(
    "I am Md. Rafi Ullah (ID: 2251081204), a student of Batch 60-B, Department of Computer Science and Engineering, Uttara University. Under the supervision of Md. Injamul Islam, Lecturer and Coordinator, Department of CSE, I am developing an AI-based smart farming system. The system will analyze soil and crop-related information to assess land suitability, recommend appropriate crops, and provide basic cultivation guidance."
)

add("For training and validating the crop-suitability model and testing its integration with the application, I respectfully request access to:", after=3)

items = [
    "Crop datasets: crop name and variety, suitable soil type, season, location, irrigation and fertilizer requirements, cultivation conditions, and actual yield.",
    "Soil datasets: moisture, temperature, pH, EC, NPK, organic matter, soil type, and corresponding crop suitability or yield.",
    "Temporary supervised access to multi-parameter soil-testing equipment capable of measuring moisture, temperature, pH, EC, and NPK, for testing real soil samples and application integration.",
]
for item in items:
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.first_line_indent = Inches(-0.18)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(item)
    r.font.name = "Times New Roman"; r.font.size = Pt(11)

p = add(after=4)
p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
p.add_run("The requested resources will be used solely for academic research, model development, and prototype validation under institutional supervision. Any conditions concerning data access, confidentiality, and equipment use will be followed. Sher-e-Bangla Agricultural University and the concerned departments will be duly acknowledged in the project report and any future publications.")

add("I therefore request your kind permission and cooperation regarding the aforementioned datasets and equipment.", after=6)
add("Yours faithfully,", after=2)
add("Md. Rafi Ullah", bold=True, after=0)
add("ID: 2251081204 | Batch: 60-B", after=0)
add("Department of CSE, Uttara University", after=0)
add("Mobile: 01606079896", after=7)

add("Supervised by:", bold=True, after=14)
add("Md. Injamul Islam", after=0)
add("Lecturer and Coordinator", after=0)
add("Department of CSE, Uttara University", after=6)

add("Recommended and forwarded by:", bold=True, after=14)
add("Dr. S. M. Nazmus Sadat", after=0)
add("Assistant Professor and Chairman", after=0)
add("Department of CSE, Uttara University", after=0)

doc.save(OUT)
print(OUT)
